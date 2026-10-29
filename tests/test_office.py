"""Word・Excel・PowerPoint・PDF の作成。"""

import io

import pytest

from lao.documents import extract_text
from lao.office import OfficeError, parse_markdown, render

MD = """# 月次報告
## 概要
今月は **売上が増加** しました。
- 新規顧客 12 社
  - うち大口 3 社
1. 次の手
> 速報値です
| 項目 | 値 |
|---|---|
| 売上 | 1,200 |
"""


def test_parse_markdown():
    kinds = [(b.kind, b.level) for b in parse_markdown(MD)]
    assert kinds == [("h1", 0), ("h2", 0), ("p", 0), ("ul", 0), ("ul", 1), ("ol", 0), ("quote", 0), ("table", 0)]
    assert parse_markdown(MD)[-1].rows == [["項目", "値"], ["売上", "1,200"]]


@pytest.mark.parametrize("fmt", ["docx", "pdf"])
def test_documents_roundtrip(tmp_path, fmt):
    p = tmp_path / f"a.{fmt}"
    p.write_bytes(render(fmt, "月次報告", MD))
    text = extract_text(p)
    for word in ("月次報告", "売上が増加", "うち大口", "1,200"):
        assert word in text


def test_xlsx_sheets_numbers_formulas():
    from openpyxl import load_workbook
    data = render("xlsx", "集計", "## 売上\n| 月 | 金額 |\n|---|---|\n| 1月 | 1,200 |\n| 合計 | =SUM(B2:B2) |\n"
                                  "## メモ\n注意,速報")
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["売上", "メモ"]
    ws = wb["売上"]
    assert ws["B2"].value == 1200 and ws["B3"].value == "=SUM(B2:B2)" and ws["A1"].font.bold
    assert wb["メモ"]["B1"].value == "速報"


def test_pptx_slides_and_notes(tmp_path):
    from pptx import Presentation
    data = render("pptx", "報告", "# 報告\n2026年\n---\n# 概要\n- 一\n- 二\nノート：補足\n---\n# 表\n| a | b |\n|---|---|\n| 1 | 2 |")
    prs = Presentation(io.BytesIO(data))
    assert [s.shapes.title.text for s in prs.slides] == ["報告", "概要", "表"]
    assert prs.slides[1].notes_slide.notes_text_frame.text == "補足"
    p = tmp_path / "a.pptx"
    p.write_bytes(data)
    assert "1 | 2" in extract_text(p)


def test_render_rejects():
    with pytest.raises(OfficeError):
        render("exe", "x", "y")
    with pytest.raises(OfficeError):
        render("docx", "x", "  ")
