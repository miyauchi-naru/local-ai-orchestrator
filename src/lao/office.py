"""Word・Excel・PowerPoint・PDF を作る。

AI は中身を Markdown で書き、ここで各形式に変換する（AI にプログラムを書かせて実行することはしない）。

- docx / pdf：見出し（#）、箇条書き（- / 1.）、表（| a | b |）、引用（>）、太字（**x**）、段落
- xlsx：「## シート名」でシートを分け、表（| a | b |）かカンマ区切りの行を並べる。= で始まるセルは式
- pptx：「---」の行でスライドを分ける。各スライドの最初の # 行が題、箇条書きが本文、「ノート：」以降は発表者ノート
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

FORMATS = {".docx", ".xlsx", ".pptx", ".pdf"}
MAX_CHARS = 200_000


class OfficeError(ValueError):
    pass


@dataclass
class Block:
    kind: str  # h1..h6 / p / ul / ol / table / quote
    text: str = ""
    rows: list[list[str]] = field(default_factory=list)
    level: int = 0


_TABLE_SEP = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?$")


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def parse_markdown(text: str) -> list[Block]:
    blocks: list[Block] = []
    para: list[str] = []
    lines = text.replace("\r\n", "\n").split("\n")

    def flush():
        if para:
            blocks.append(Block("p", " ".join(s.strip() for s in para)))
            para.clear()

    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s or s.startswith("```"):
            flush()
        elif m := re.match(r"^(#{1,6})\s+(.*)$", s):
            flush()
            blocks.append(Block(f"h{len(m.group(1))}", m.group(2).strip()))
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not _TABLE_SEP.match(lines[i].strip()):
                    rows.append(_cells(lines[i]))
                i += 1
            blocks.append(Block("table", rows=rows))
            continue
        elif m := re.match(r"^(\s*)[-*・]\s+(.*)$", line):
            flush()
            blocks.append(Block("ul", m.group(2).strip(), level=len(m.group(1).expandtabs(2)) // 2))
        elif m := re.match(r"^(\s*)\d+[.)]\s+(.*)$", line):
            flush()
            blocks.append(Block("ol", m.group(2).strip(), level=len(m.group(1).expandtabs(2)) // 2))
        elif s.startswith(">"):
            flush()
            blocks.append(Block("quote", s.lstrip("> ").strip()))
        else:
            para.append(s)
        i += 1
    flush()
    return blocks


_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _runs(text: str) -> list[tuple[str, bool]]:
    """**太字** を (文字, 太字か) に分ける。"""
    out, pos = [], 0
    for m in _BOLD.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], False))
        out.append((m.group(1), True))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False))
    return out or [("", False)]


def _plain(text: str) -> str:
    return _BOLD.sub(r"\1", text)


# ---------------------------------------------------------------- Word
def to_docx(title: str, text: str) -> bytes:
    from docx import Document
    doc = Document()
    if title:
        doc.core_properties.title = title
    for b in parse_markdown(text):
        if b.kind.startswith("h"):
            doc.add_heading(_plain(b.text), level=min(int(b.kind[1]), 4))
        elif b.kind == "table":
            if not b.rows:
                continue
            width = max(len(r) for r in b.rows)
            t = doc.add_table(rows=len(b.rows), cols=width)
            t.style = "Table Grid"
            for r, row in enumerate(b.rows):
                for c in range(width):
                    cell = t.cell(r, c)
                    cell.text = ""
                    for s, bold in _runs(row[c] if c < len(row) else ""):
                        cell.paragraphs[0].add_run(s).bold = bold or r == 0
        else:
            style = {"ul": "List Bullet", "ol": "List Number", "quote": "Quote"}.get(b.kind)
            if style and b.kind in ("ul", "ol") and b.level:
                style += f" {min(b.level + 1, 3)}"
            p = doc.add_paragraph(style=style)
            for s, bold in _runs(b.text):
                p.add_run(s).bold = bold
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- Excel
def _value(s: str):
    s = s.strip()
    if s.startswith("="):
        return s
    t = s.replace(",", "")
    try:
        return int(t) if re.fullmatch(r"-?\d+", t) else float(t) if re.fullmatch(r"-?\d+\.\d+", t) else s
    except ValueError:
        return s


def to_xlsx(title: str, text: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    sheets: list[tuple[str, list[list[str]]]] = []
    name, rows = "Sheet1", []
    for line in text.replace("\r\n", "\n").split("\n"):
        s = line.strip()
        if m := re.match(r"^#{1,6}\s+(.*)$", s):
            if rows:
                sheets.append((name, rows))
            name, rows = m.group(1).strip(), []
        elif s.startswith("|"):
            if not _TABLE_SEP.match(s):
                rows.append(_cells(s))
        elif s and not s.startswith("```"):
            rows.append([c.strip() for c in s.split(",")])
    if rows or not sheets:
        sheets.append((name, rows))
    used = set()
    for name, rows in sheets:
        safe = re.sub(r"[\[\]:*?/\\]", "_", name)[:31] or "Sheet"
        while safe in used:
            safe = safe[:28] + f"_{len(used)}"
        used.add(safe)
        ws = wb.create_sheet(safe)
        for r in rows:
            ws.append([_value(c) for c in r])
        if rows:
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for col in ws.columns:
                width = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(8, width * 1.6), 60)
    if title:
        wb.properties.title = title
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PowerPoint
def to_pptx(title: str, text: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Pt
    prs = Presentation()
    slides = re.split(r"(?m)^\s*---+\s*$", text.replace("\r\n", "\n"))
    for i, chunk in enumerate(s for s in slides if s.strip()):
        body, notes = chunk, ""
        m = re.search(r"(?m)^\s*(ノート|notes?)\s*[:：]", chunk, re.I)
        if m:
            body, notes = chunk[:m.start()], chunk[m.end():].strip()
        blocks = parse_markdown(body)
        head = next((b for b in blocks if b.kind.startswith("h")), None)
        rest = [b for b in blocks if b is not head]
        layout = prs.slide_layouts[0 if i == 0 and len(rest) <= 1 and not any(b.kind == "table" for b in rest) else 1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = _plain(head.text) if head else (title if i == 0 else "")
        ph = slide.placeholders[1]
        tables = [b for b in rest if b.kind == "table" and b.rows]
        lines = [b for b in rest if b.kind != "table"]
        if layout == prs.slide_layouts[0]:
            ph.text = _plain(lines[0].text) if lines else ""
        else:
            tf = ph.text_frame
            tf.clear()
            for n, b in enumerate(lines):
                p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
                p.level = min(b.level + (1 if b.kind.startswith("h") else 0), 4)
                for s, bold in _runs(b.text):
                    run = p.add_run()
                    run.text = s
                    run.font.bold = bold or b.kind.startswith("h")
            if tables and not lines:
                ph._element.getparent().remove(ph._element)
            for t in tables[:1]:
                width = max(len(r) for r in t.rows)
                top = ph.top + ph.height // 2 if lines else prs.slide_height // 4
                shape = slide.shapes.add_table(len(t.rows), width, prs.slide_width // 12, top,
                                               prs.slide_width * 10 // 12, Pt(28) * len(t.rows))
                for r, row in enumerate(t.rows):
                    for c in range(width):
                        shape.table.cell(r, c).text = _plain(row[c] if c < len(row) else "")
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    if title:
        prs.core_properties.title = title
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF
# 日本語フォント：PC にある TrueType を埋め込む（使う文字だけ）。見つからなければ PDF 標準の CID フォント（埋め込まず、表示側のフォントで描く）
PDF_FONTS = [
    ("C:/Windows/Fonts/YuGothM.ttc", 0), ("C:/Windows/Fonts/meiryo.ttc", 0), ("C:/Windows/Fonts/msgothic.ttc", 0),
    ("/usr/share/fonts/opentype/ipafont-gothic/ipag.ttf", 0), ("/usr/share/fonts/truetype/fonts-japanese-gothic.ttf", 0),
]


def _pdf_font(pdfmetrics, UnicodeCIDFont) -> str:
    import os
    from pathlib import Path

    from reportlab.pdfbase.ttfonts import TTFont
    if "LaoJP" in pdfmetrics.getRegisteredFontNames():
        return "LaoJP"
    extra = os.environ.get("LAO_PDF_FONT")
    for path, index in ([(extra, 0)] if extra else []) + PDF_FONTS:
        if Path(path).is_file():
            try:
                pdfmetrics.registerFont(TTFont("LaoJP", path, subfontIndex=index))
                return "LaoJP"
            except Exception:  # CFF の OpenType など reportlab が読めない形式
                continue
    if "HeiseiKakuGo-W5" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont("HeiseiKakuGo-W5"))
    return "HeiseiKakuGo-W5"


def to_pdf(title: str, text: str) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = _pdf_font(pdfmetrics, UnicodeCIDFont)
    base = ParagraphStyle("base", fontName=font, fontSize=10.5, leading=17, wordWrap="CJK")
    styles = {f"h{n}": ParagraphStyle(f"h{n}", parent=base, fontSize=size, leading=size * 1.5, spaceBefore=8,
                                       spaceAfter=4) for n, size in zip(range(1, 7), (20, 16, 14, 12, 11, 10.5))}
    quote = ParagraphStyle("quote", parent=base, leftIndent=8 * mm, textColor=colors.HexColor("#555555"))

    def markup(s: str) -> str:
        return "".join(f"<b>{escape(t)}</b>" if bold else escape(t) for t, bold in _runs(s))

    story = []
    if title:
        story.append(Paragraph(escape(title), styles["h1"]))
    blocks = parse_markdown(text)
    i = 0
    while i < len(blocks):
        b = blocks[i]
        if b.kind in ("ul", "ol"):
            items = []
            while i < len(blocks) and blocks[i].kind == b.kind:
                items.append(ListItem(Paragraph(markup(blocks[i].text), base), leftIndent=(6 + blocks[i].level * 6) * mm))
                i += 1
            story.append(ListFlowable(items, bulletType="bullet" if b.kind == "ul" else "1", bulletFontName=font,
                                      start="•" if b.kind == "ul" else None))
            continue
        if b.kind.startswith("h"):
            if not (i == 0 and title and _plain(b.text) == title):
                story.append(Paragraph(markup(b.text), styles[b.kind]))
        elif b.kind == "table" and b.rows:
            width = max(len(r) for r in b.rows)
            data = [[Paragraph(markup(r[c] if c < len(r) else ""), base) for c in range(width)] for r in b.rows]
            t = Table(data, repeatRows=1, colWidths=[(A4[0] - 40 * mm) / width] * width)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story += [t, Spacer(1, 4 * mm)]
        elif b.kind == "quote":
            story.append(Paragraph(markup(b.text), quote))
        else:
            story += [Paragraph(markup(b.text), base), Spacer(1, 2 * mm)]
        i += 1
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=20 * mm,
                      bottomMargin=20 * mm, title=title or "", author="").build(story or [Spacer(1, 1)])
    return buf.getvalue()


def render(fmt: str, title: str, text: str) -> bytes:
    fmt = fmt.lower() if fmt.startswith(".") else "." + fmt.lower()
    if fmt not in FORMATS:
        raise OfficeError(f"形式は {sorted(FORMATS)} のいずれかです")
    if not text.strip():
        raise OfficeError("中身が空です")
    if len(text) > MAX_CHARS:
        raise OfficeError(f"中身は {MAX_CHARS} 文字までです")
    try:
        return {".docx": to_docx, ".xlsx": to_xlsx, ".pptx": to_pptx, ".pdf": to_pdf}[fmt](title, text)
    except ImportError as e:  # pragma: no cover
        raise OfficeError(f"{fmt} を作る部品がありません：{e}") from e
