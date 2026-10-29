"""文書の添付（テキスト・PDF）から本文を取り出す。中身はデータとして扱い、命令としては扱わない。"""

from __future__ import annotations

from pathlib import Path

TEXT_SUFFIXES = {".txt", ".md", ".csv", ".tsv", ".json", ".jsonl", ".yaml", ".yml", ".toml", ".xml", ".html",
                 ".htm", ".css", ".js", ".ts", ".py", ".ps1", ".sh", ".bat", ".c", ".cpp", ".h", ".cs", ".java",
                 ".go", ".rs", ".sql", ".log", ".ini", ".srt", ".vtt"}
OFFICE_SUFFIXES = {".docx", ".xlsx", ".pptx"}
DOC_SUFFIXES = TEXT_SUFFIXES | OFFICE_SUFFIXES | {".pdf"}
MAX_CHARS = 100_000


class DocumentError(ValueError):
    pass


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as e:  # pragma: no cover
            raise DocumentError("PDF を読むには pypdf が必要です（pip install pypdf）") from e
        try:
            reader = PdfReader(str(path))
            pages = [p.extract_text() or "" for p in reader.pages]
        except Exception as e:  # pypdf は壊れたファイルで様々な例外を出す
            raise DocumentError(f"PDF を読めませんでした：{e}") from e
        text = "\n\n".join(f"[{i} ページ]\n{t.strip()}" for i, t in enumerate(pages, 1) if t.strip())
        if not text:
            raise DocumentError("PDF から文字を取り出せませんでした（画像だけの PDF は画像として添付してください）")
        return text
    if suffix in OFFICE_SUFFIXES:
        try:
            text = _office_text(path, suffix)
        except Exception as e:  # 壊れたファイルで各ライブラリが様々な例外を出す
            raise DocumentError(f"{suffix} を読めませんでした：{e}") from e
        if not text.strip():
            raise DocumentError(f"{suffix} から文字を取り出せませんでした")
        return text
    if suffix in TEXT_SUFFIXES:
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp932"):
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                continue
        return raw.decode("utf-8", errors="replace")
    raise DocumentError(f"{suffix} は文書として読めません")


def _office_text(path: Path, suffix: str) -> str:
    if suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        out = [p.text for p in doc.paragraphs if p.text.strip()]
        for i, t in enumerate(doc.tables, 1):
            out.append(f"[表 {i}]")
            out += [" | ".join(c.text.strip() for c in row.cells) for row in t.rows]
        return "\n".join(out)
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"[シート {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    out.append(",".join("" if v is None else str(v) for v in row))
        return "\n".join(out)
    from pptx import Presentation
    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"[スライド {i}]")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append("（ノート）" + slide.notes_slide.notes_text_frame.text)
    return "\n".join(out)


def chunks(text: str, size: int = 800) -> list[str]:
    paras, out, cur = text.split("\n"), [], ""
    for p in paras:
        while len(p) > size:
            if cur:
                out.append(cur)
                cur = ""
            out.append(p[:size])
            p = p[size:]
        if len(cur) + len(p) + 1 > size and cur:
            out.append(cur)
            cur = ""
        cur = f"{cur}\n{p}" if cur else p
    if cur.strip():
        out.append(cur)
    return out


def _terms(text: str) -> list[str]:
    """日本語は単語の区切りがないため、文字の 2-gram と英数字の単語で数える。"""
    import re
    text = text.lower()
    words = re.findall(r"[a-z0-9_]+", text)
    cjk = re.sub(r"[a-z0-9_\s\W]+", " ", text)
    grams = [w[i:i + 2] for w in cjk.split() for i in range(max(1, len(w) - 1))]
    return words + grams


def best_chunks(docs: list[tuple[str, str]], query: str, budget: int) -> list[tuple[str, str]]:
    """資料（名前, 本文）から、query に近い部分を budget 文字まで選ぶ（BM25）。"""
    import math
    from collections import Counter
    pieces = [(name, c) for name, text in docs for c in chunks(text)]
    if not pieces:
        return []
    tfs = [Counter(_terms(c)) for _, c in pieces]
    df = Counter(t for tf in tfs for t in tf)
    avg = sum(sum(tf.values()) for tf in tfs) / len(tfs) or 1
    q = set(_terms(query))
    n = len(pieces)

    def score(tf):
        dl = sum(tf.values()) or 1
        s = 0.0
        for t in q:
            if t in tf:
                idf = math.log(1 + (n - df[t] + 0.5) / (df[t] + 0.5))
                s += idf * tf[t] * 2.2 / (tf[t] + 1.2 * (0.25 + 0.75 * dl / avg))
        return s

    order = sorted(range(n), key=lambda i: score(tfs[i]), reverse=True)
    out, used = [], 0
    for i in order:
        if used + len(pieces[i][1]) > budget:
            continue
        out.append(pieces[i])
        used += len(pieces[i][1])
    return out


def clip(text: str, limit: int = MAX_CHARS) -> tuple[str, bool]:
    return (text, False) if len(text) <= limit else (text[:limit], True)
